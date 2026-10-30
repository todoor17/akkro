#!/usr/bin/env python3
"""Reads every message of the AKKRO Outlook PST into one JSON file, for genereaza_emailuri.py.

Runs inside the akkro-pst image (Dockerfile.pst), which has libpff and the text extractors:

    docker build -t akkro-pst -f Dockerfile.pst .
    docker run --rm --user "$(id -u):$(id -g)" -v "<folder with the pst>":/in:ro -v "$PWD":/work akkro-pst \
        python3 extrage_pst.py /in/exportAI.pst

Outputs:
  - date/mesaje.json: one entry per message (all folders), with headers, recipients, body as text and the
    attachments (name, type, size, extracted text or the reason there is none);
  - output/atasamente/<message id>/<file>: the attachments, except inline images (signatures, logos).

Both contain client data: internal only.
"""

import argparse
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import pypff
from bs4 import BeautifulSoup

HERE = Path(__file__).resolve().parent
# MAPI property tags
PR_MESSAGE_CLASS = 0x001A
PR_CONVERSATION_TOPIC = 0x0070
PR_CONVERSATION_INDEX = 0x0071
PR_SENDER_NAME = 0x0C1A
PR_SENDER_ADDRTYPE = 0x0C1E
PR_SENDER_EMAIL = 0x0C1F
PR_SENDER_SMTP = 0x5D01
PR_SENT_REPRESENTING_EMAIL = 0x0065
PR_INTERNET_MESSAGE_ID = 0x1035
PR_INTERNET_REFERENCES = 0x1039
PR_IN_REPLY_TO = 0x1042
PR_RECIPIENT_TYPE = 0x0C15
PR_DISPLAY_NAME = 0x3001
PR_EMAIL_ADDRESS = 0x3003
PR_SMTP_ADDRESS = 0x39FE
PR_ATTACH_FILENAME = 0x3704
PR_ATTACH_LONG_FILENAME = 0x3707
PR_ATTACH_CONTENT_ID = 0x3712
RECIPIENT_FIELD = {1: "to", 2: "cc", 3: "bcc"}
STRING_TYPES = (0x001E, 0x001F)
# headers kept for classifying automatic messages (newsletters, auto-replies, delivery reports)
KEPT_HEADERS = (
    "auto-submitted",
    "precedence",
    "list-unsubscribe",
    "list-id",
    "x-auto-response-suppress",
    "x-autoreply",
    "x-autorespond",
    "thread-topic",
    "thread-index",
    "message-id",
    "in-reply-to",
    "references",
    "x-spam-flag",
    "x-bitdefender-spam",
)
TEXT_LIMIT = 60000  # characters kept from one attachment
IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff", "webp", "emz", "wmz"}


def props(item):
    """{tag: entry} for the first occurrence of each property of a message, recipient or attachment."""
    found = {}
    for rs in item.record_sets:
        for e in rs.entries:
            found.setdefault(e.entry_type, e)
    return found


def as_text(entry):
    if entry is None or entry.data is None:
        return ""
    if entry.value_type in STRING_TYPES:
        try:
            return entry.data_as_string.strip("\x00").strip()
        except Exception:
            return entry.data.decode("utf-16-le" if entry.value_type == 0x001F else "latin-1", "ignore").strip("\x00")
    return ""


def as_int(entry):
    try:
        return entry.data_as_integer if entry is not None else None
    except Exception:
        return None


def parse_headers(raw):
    headers = {}
    for name, value in re.findall(r"^([A-Za-z0-9-]+):[ \t]*(.*(?:\r?\n[ \t].*)*)", raw or "", re.M):
        key = name.lower()
        if key in KEPT_HEADERS and key not in headers:
            headers[key] = re.sub(r"\s+", " ", value).strip()
    return headers


def html_to_text(raw):
    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["style", "script", "head", "title"]):
        tag.decompose()
    for br in soup.find_all("br"):
        br.replace_with("\n")
    for block in soup.find_all(["p", "div", "tr", "li", "h1", "h2", "h3", "h4", "table", "hr", "blockquote"]):
        block.insert_before("\n")
        block.insert_after("\n")
    for cell in soup.find_all(["td", "th"]):
        cell.insert_after(" ")
    return soup.get_text()


def decode(raw):
    for enc in ("utf-8", "cp1250", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "ignore")


def tidy(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ").replace("​", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def message_body(m):
    html = m.html_body
    if html:
        return tidy(html_to_text(decode(html)))
    if m.plain_text_body:
        return tidy(decode(m.plain_text_body))
    return ""


def address(entry_props):
    smtp = as_text(entry_props.get(PR_SMTP_ADDRESS))
    email = as_text(entry_props.get(PR_EMAIL_ADDRESS))
    return (smtp or (email if "@" in email else "")).lower()


def recipients(m):
    out = []
    try:
        table = m.recipients
    except Exception:
        return out
    if table is None:
        return out
    for rs in table.record_sets:
        p = {e.entry_type: e for e in rs.entries}
        out.append(
            {
                "field": RECIPIENT_FIELD.get(as_int(p.get(PR_RECIPIENT_TYPE)), "to"),
                "name": as_text(p.get(PR_DISPLAY_NAME)).strip("'\" "),
                "email": address(p),
            }
        )
    return out


def run_text(cmd, data, suffix):
    with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:
        tmp.write(data)
        tmp.flush()
        res = subprocess.run(
            cmd + [tmp.name] + (["-"] if cmd[0] == "pdftotext" else []), capture_output=True, timeout=120
        )
    return decode(res.stdout)


def office_text(data, ext):
    import io

    if ext == "docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def attachment_text(data, ext):
    """(text, reason): text extracted from an attachment, or why there is none."""
    try:
        if ext == "pdf":
            text = run_text(["pdftotext", "-layout"], data, ".pdf")
            return (text, "") if text.strip() else ("", "PDF fără text (scanat)")
        if ext in ("docx", "xlsx", "xlsm"):
            return office_text(data, ext), ""
        if ext == "doc":
            return run_text(["catdoc", "-w"], data, ".doc"), ""
        if ext == "xls":
            return run_text(["xls2csv"], data, ".xls"), ""
        if ext in ("txt", "csv"):
            return decode(data), ""
        if ext in ("html", "htm"):
            return tidy(html_to_text(decode(data))), ""
    except Exception as exc:  # a broken file must not stop the extraction
        return "", f"eroare la citire: {exc.__class__.__name__}"
    if ext in IMAGE_EXT:
        return "", "imagine"
    return "", f"tip necitit (.{ext})" if ext else "fără extensie"


def safe_name(name):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .") or "atasament"


def attachments(m, msg_id, out_dir):
    out = []
    for i in range(m.number_of_attachments):
        try:
            att = m.get_attachment(i)
        except Exception as exc:
            out.append(
                {
                    "name": f"atașament {i + 1}",
                    "ext": "",
                    "size": 0,
                    "inline": False,
                    "path": "",
                    "text": "",
                    "reason": f"eroare la citire: {exc.__class__.__name__}",
                }
            )
            continue
        p = props(att)
        name = as_text(p.get(PR_ATTACH_LONG_FILENAME)) or as_text(p.get(PR_ATTACH_FILENAME))
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        inline = bool(as_text(p.get(PR_ATTACH_CONTENT_ID))) and ext in IMAGE_EXT
        entry = {
            "name": name or f"atașament {i + 1}",
            "ext": ext,
            "size": 0,
            "inline": inline,
            "path": "",
            "text": "",
            "reason": "",
        }
        try:
            size = att.size or 0
            data = att.read_buffer(size) if size else b""
        except Exception:
            # an attached Outlook message: keep its subject and body as the text
            data = b""
            try:
                item = att.get_item()
                entry["text"] = tidy(f"{item.subject or ''}\n\n{message_body(item)}")
                entry["ext"] = entry["ext"] or "msg"
                entry["reason"] = "" if entry["text"] else "mesaj atașat gol"
            except Exception as exc:
                entry["reason"] = f"eroare la citire: {exc.__class__.__name__}"
            out.append(entry)
            continue
        entry["size"] = len(data)
        if not inline and data:
            folder = out_dir / msg_id
            folder.mkdir(parents=True, exist_ok=True)
            target = folder / safe_name(entry["name"])
            n = 2
            while target.exists():
                target = folder / f"{n}_{safe_name(entry['name'])}"
                n += 1
            target.write_bytes(data)
            entry["path"] = target.relative_to(out_dir.parent).as_posix()
            text, reason = attachment_text(data, ext)
            entry["text"] = tidy(text)[:TEXT_LIMIT]
            entry["reason"] = reason or ("" if entry["text"] else "fără text")
        elif inline:
            entry["reason"] = "imagine în corpul mesajului"
        out.append(entry)
    return out


def iso(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ") if dt else ""


def read_message(m, folder_path, out_dir):
    p = props(m)
    msg_id = f"m{m.identifier}"
    raw_headers = m.transport_headers or ""
    headers = parse_headers(raw_headers)
    sender_type = as_text(p.get(PR_SENDER_ADDRTYPE))
    sender = as_text(p.get(PR_SENDER_SMTP)) or as_text(p.get(PR_SENDER_EMAIL))
    if "@" not in sender:
        sender = as_text(p.get(PR_SENT_REPRESENTING_EMAIL)) if "@" in as_text(p.get(PR_SENT_REPRESENTING_EMAIL)) else ""
    conv_index = p.get(PR_CONVERSATION_INDEX)
    sent = m.client_submit_time
    received = m.delivery_time
    return {
        "id": msg_id,
        "folder": folder_path,
        "class": as_text(p.get(PR_MESSAGE_CLASS)),
        "subject": (m.subject or "").strip(),
        "date": iso(received or sent or m.creation_time),
        "sent": iso(sent),
        "from_name": (m.sender_name or as_text(p.get(PR_SENDER_NAME))).strip(),
        "from_email": sender.lower(),
        "from_type": sender_type,
        "recipients": recipients(m),
        "message_id": as_text(p.get(PR_INTERNET_MESSAGE_ID)) or headers.get("message-id", ""),
        "in_reply_to": as_text(p.get(PR_IN_REPLY_TO)) or headers.get("in-reply-to", ""),
        "references": as_text(p.get(PR_INTERNET_REFERENCES)) or headers.get("references", ""),
        "conversation_index": conv_index.data.hex() if conv_index is not None and conv_index.data else "",
        "conversation_topic": as_text(p.get(PR_CONVERSATION_TOPIC)),
        "headers": headers,
        "body": message_body(m),
        "attachments": attachments(m, msg_id, out_dir),
    }


def walk(folder, path):
    yield folder, path
    for k in range(folder.number_of_sub_folders):
        sub = folder.get_sub_folder(k)
        yield from walk(sub, f"{path}/{sub.name}" if path else (sub.name or ""))


def main():
    ap = argparse.ArgumentParser(description="Extrage toate mesajele din PST în date/mesaje.json.")
    ap.add_argument("pst", help="fișierul Outlook .pst")
    ap.add_argument("--out", default=str(HERE / "date" / "mesaje.json"), help="fișierul JSON scris")
    ap.add_argument("--atasamente", default=str(HERE / "output" / "atasamente"), help="unde se salvează atașamentele")
    args = ap.parse_args()

    pst = pypff.file()
    pst.open(args.pst)
    out_dir = Path(args.atasamente)
    messages, errors = [], []
    for folder, path in walk(pst.root_folder, ""):
        # the top-level "Top of Outlook data file/" prefix carries no information
        short = path.split("/", 1)[1] if path.startswith("Top of Outlook data file/") else path
        for i in range(folder.number_of_sub_messages):
            try:
                messages.append(read_message(folder.get_sub_message(i), short, out_dir))
            except Exception as exc:
                # kept as a stub, so the report still adds up to the number of messages in the PST
                errors.append(f"{short} #{i}: {exc.__class__.__name__}: {exc}")
                messages.append({"id": f"err{len(errors)}", "folder": short, "error": str(exc)})
        if folder.number_of_sub_messages:
            print(f"  {short}: {folder.number_of_sub_messages} mesaje", file=sys.stderr)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(
        json.dumps({"pst": Path(args.pst).name, "messages": messages, "errors": errors}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    print(f"Gata: {len(messages)} mesaje în {out}, {len(errors)} erori.", file=sys.stderr)
    for e in errors:
        print(f"  ! {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
